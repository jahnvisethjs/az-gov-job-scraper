from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io
from typing import Dict

class ResumeGenerator:
    """Generates a tailored Word document resume (.docx) based on AI advice."""
    
    @staticmethod
    def generate_docx(profile: Dict, advice: Dict) -> bytes:
        """
        Creates a formatted .docx file with the original resume content + tailored suggestions.
        
        Args:
            profile: User profile dictionary containing original resume text/parsed data
            advice: Dictionary with tailoring advice (strengths, keywords, etc.)
            
        Returns:
            bytes: The .docx file content as bytes
        """
        doc = Document()
        
        # ── Setup Margins ──
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)
            
        # ── Header (Name & Contact) ──
        name = profile.get("name") or "Your Name"
        heading = doc.add_heading(name, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Add basic info
        contact_para = doc.add_paragraph()
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        degree = profile.get("degree", "")
        if degree:
            contact_para.add_run(f"Education: {degree} | ")
        contact_para.add_run("Tailored AI Resume")
            
        doc.add_paragraph() # Spacing
        
        # ── Professional Summary (AI Tailored) ──
        doc.add_heading("Professional Summary", level=1)
        summary = doc.add_paragraph()
        
        # We craft a quick custom summary based on their strengths and the required keywords
        strengths = advice.get("strengths", [])
        keywords = advice.get("keywords", [])
        
        # Helper to join list naturally
        def join_list(lst):
            if not lst: return ""
            if len(lst) == 1: return lst[0]
            if len(lst) == 2: return f"{lst[0]} and {lst[1]}"
            return f"{', '.join(lst[:-1])}, and {lst[-1]}"
            
        summary_text = ""
        if strengths:
            summary_text += f"Dedicated professional with strong background in {join_list(strengths[:3])}. "
        if keywords:
            summary_text += f"Skilled in areas including {join_list(keywords[:4])}. "
        
        if summary_text:
            summary.add_run(summary_text)
        else:
            summary.add_run("Highly motivated professional looking to leverage skills and experience to contribute to team success.")
            
        # ── Core Competencies / Skills (AI Augmented) ──
        doc.add_heading("Core Competencies & Skills", level=1)
        
        # Combine original skills + new AI suggested keywords
        original_skills = profile.get("resume_parsed", {}).get("skills", [])
        all_skills = list(set(original_skills + keywords)) # Remove duplicates
        
        # Format as a bulleted list or comma separated
        skills_para = doc.add_paragraph(style='List Bullet')
        for skill in all_skills[:12]: # Cap at 12 skills
            doc.add_paragraph(skill, style='List Bullet')
            
        doc.add_paragraph() # Spacing
        
        # ── Experience ──
        doc.add_heading("Professional Experience", level=1)
        
        experience = profile.get("resume_parsed", {}).get("experience", [])
        if experience:
            for exp in experience:
                if isinstance(exp, dict):
                    # Structured experience
                    title = exp.get("title", "Position")
                    company = exp.get("company", "Company")
                    date = exp.get("date", "Dates")
                    
                    p = doc.add_paragraph()
                    p.add_run(title).bold = True
                    p.add_run(f" | {company}").italic = True
                    p.add_run(f" | {date}")
                    
                    desc = exp.get("description", "")
                    if desc:
                        doc.add_paragraph(desc, style='List Bullet')
                else:
                    # Unstructured string
                    doc.add_paragraph(str(exp), style='List Bullet')
        else:
            # Fallback to pure text if we don't have structured experience
            resume_text = profile.get("resume_text", "")
            if resume_text:
                # Add a snippet of the raw text so it's not empty
                doc.add_paragraph("(Experience extracted from original document:)")
                # Just add first 1000 chars of raw text for placeholder
                doc.add_paragraph(resume_text[:1000] + "...")
            else:
                doc.add_paragraph("[Add your professional experience here]")
                
        # ── AI Suggestions Note (For the user) ──
        doc.add_page_break()
        doc.add_heading("AI Tailoring Advice (Remove before applying!)", level=1)
        
        doc.add_paragraph("This resume was generated to highlight your qualifications for this specific job. To maximize your chances, consider these AI recommendations:", style='Intense Quote')
        
        doc.add_heading("Address these gaps:", level=2)
        for gap in advice.get("skill_gaps", []):
            doc.add_paragraph(f"• {gap}")
            
        doc.add_heading("General Improvements:", level=2)
        for imp in advice.get("improvements", []):
            doc.add_paragraph(f"• {imp}")
            
        # Save to memory buffer
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        return buffer.getvalue()
